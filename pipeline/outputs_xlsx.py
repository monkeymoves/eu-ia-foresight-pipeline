"""manifest.json + mentions.json (+ ai_drafts.json) -> data/out/dataset.xlsx

Single sheet, one row per impact assessment: the research team records one
judgement per document here; per-mention evidence lives on the review page.
Links are absolute URLs against config.PUBLIC_BASE_URL so they work from
Google Sheets, Excel, and anywhere else. See SPEC.md for the contract.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline import config

LINK_FONT = Font(color="0B5FA5", underline="single")

# --- Scope heuristic -------------------------------------------------------

_COUNCIL_DIRECTIVE_RE = re.compile(r"\bcouncil directive\b", re.IGNORECASE)
_DIRECTIVE_RE = re.compile(r"\bdirective\b", re.IGNORECASE)
_REGULATION_RE = re.compile(r"\bregulation\b", re.IGNORECASE)
_DECISION_RE = re.compile(r"\bdecision\b", re.IGNORECASE)
_AMENDING_RE = re.compile(r"\bamending\b", re.IGNORECASE)


def derive_scope(title: str) -> str:
    """Best-effort enum classification of the legal act from the IA title."""
    types: set[str] = set()
    if _COUNCIL_DIRECTIVE_RE.search(title):
        types.add("council directive")
    elif _DIRECTIVE_RE.search(title):
        types.add("directive")
    if _REGULATION_RE.search(title):
        types.add("regulation")
    if _DECISION_RE.search(title):
        types.add("decision")

    if len(types) > 1:
        return "multiple acts"
    if not types:
        return "other"

    only = next(iter(types))
    amending = bool(_AMENDING_RE.search(title))
    if only == "council directive":
        return "council directive"
    if only == "decision":
        return "decision"
    if only == "directive":
        return "amending directive" if amending else "new directive"
    if only == "regulation":
        return "amending regulation" if amending else "new regulation"
    return "other"  # pragma: no cover


def load_json(path: Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def mention_pages_summary(mentions: list[dict]) -> str:
    """"part 1: p.27, p.136; part 3: p.15" - orientation only; the review
    page is where the excerpts live."""
    by_part: dict[int, list[int]] = {}
    for m in mentions:
        by_part.setdefault(m["part"], []).append(m["page"])
    bits = []
    for part in sorted(by_part):
        pages = ", ".join(f"p.{p}" for p in sorted(set(by_part[part])))
        bits.append(f"part {part}: {pages}")
    return "; ".join(bits)


HEADER_BASE = [
    "Reference",
    "Title",
    "Date",
    "Responsible service",
    "Topic",
    "AI draft: topic",
    "Scope",
    "Foresight mention",
    "Mention count",
    "Mention pages",
    "Review page",
    "Kind of foresight (final)",
    "AI draft: kind",
    "AI draft: methods",
    "AI draft: material influence",
    "AI draft: how",
    "AI draft: notes",
    "Other / notes",
]
HEADER_TAIL = ["Parts", "Flags"]


def build_workbook(manifest: dict, mentions: dict, drafts: dict | None = None) -> Workbook:
    drafts = drafts or {}
    draft_ias = drafts.get("ias", {})
    draft_topics = drafts.get("topics", {})

    related_terms = mentions.get("related_terms") or list(config.RELATED_TERMS)
    header = HEADER_BASE + [f"Related: {t}" for t in related_terms] + HEADER_TAIL

    def col(name: str) -> int:
        return header.index(name) + 1

    ias = manifest.get("ias", {})
    mention_ias = mentions.get("ias", {})
    ordered_refs = sorted(ias.keys(), key=lambda r: ias[r].get("date", ""), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "IAs"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"

    for reference in ordered_refs:
        ia = ias[reference]
        ref_safe = ia.get("ref_safe", config.ref_safe(reference))
        m = mention_ias.get(reference, {})
        mention_list = m.get("mentions", [])
        foresight_mention = m.get("foresight_mention", bool(mention_list))
        related_counts = m.get("related_term_counts", {})
        draft = draft_ias.get(reference, {})

        row = [
            reference,
            ia.get("title", ""),
            ia.get("date", ""),
            ia.get("responsible_department", ""),
            "",  # Topic - final, human-owned
            draft_topics.get(reference, ""),
            derive_scope(ia.get("title", "")),
            "YES" if foresight_mention else "no",
            len(mention_list),
            mention_pages_summary(mention_list),
            "open" if mention_list else "",
            "",  # Kind of foresight (final) - human
            draft.get("kind", ""),
            "; ".join(draft.get("methods", [])),
            draft.get("material_influence", ""),
            draft.get("influence_how", ""),
            draft.get("notes", ""),
            "",  # Other / notes - human
        ]
        row += [related_counts.get(t, 0) for t in related_terms]
        row += [len(ia.get("parts", [])), ", ".join(ia.get("flags", []))]
        ws.append(row)

        r = ws.max_row
        ref_cell = ws.cell(row=r, column=col("Reference"))
        if ia.get("detail_url"):
            ref_cell.hyperlink = ia["detail_url"]
            ref_cell.font = LINK_FONT
        if mention_list:
            review_cell = ws.cell(row=r, column=col("Review page"))
            review_cell.hyperlink = (
                f"{config.PUBLIC_BASE_URL}/out/review.html#ia-{ref_safe}"
            )
            review_cell.font = LINK_FONT

    last_row = ws.max_row

    def letter(name: str) -> str:
        return get_column_letter(col(name))

    def validate(name: str, values: list[str]) -> None:
        dv = DataValidation(
            type="list",
            formula1='"{}"'.format(",".join(values)),
            allow_blank=True,
            showDropDown=False,  # False = show the dropdown arrow (openpyxl quirk)
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter(name)}2:{letter(name)}{last_row}")

    validate("Kind of foresight (final)", config.KIND_VALUES)
    validate("AI draft: kind", config.KIND_VALUES)
    validate("AI draft: material influence", config.INFLUENCE_VALUES)

    widths = {
        "Reference": 18,
        "Title": 55,
        "Date": 12,
        "Responsible service": 14,
        "Topic": 22,
        "AI draft: topic": 22,
        "Scope": 18,
        "Foresight mention": 12,
        "Mention count": 10,
        "Mention pages": 34,
        "Review page": 12,
        "Kind of foresight (final)": 20,
        "AI draft: kind": 14,
        "AI draft: methods": 34,
        "AI draft: material influence": 14,
        "AI draft: how": 40,
        "AI draft: notes": 30,
        "Other / notes": 30,
    }
    for name, width in widths.items():
        ws.column_dimensions[letter(name)].width = width
    for term in related_terms:
        ws.column_dimensions[letter(f"Related: {term}")].width = 14
    ws.column_dimensions[letter("Parts")].width = 8
    ws.column_dimensions[letter("Flags")].width = 20

    wrap = Alignment(wrap_text=True, vertical="top")
    for name in ("Mention pages", "AI draft: how", "AI draft: notes"):
        c = col(name)
        for row_cells in ws.iter_rows(min_row=2, max_row=last_row, min_col=c, max_col=c):
            for cell in row_cells:
                cell.alignment = wrap

    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=config.MANIFEST_PATH)
    parser.add_argument("--mentions", type=Path, default=config.MENTIONS_PATH)
    parser.add_argument("--out", type=Path, default=config.OUT_DIR / "dataset.xlsx")
    parser.add_argument("--drafts", type=Path, default=config.DATA_DIR / "ai_drafts.json")
    args = parser.parse_args()

    if not args.manifest.exists():
        print(f"error: manifest not found: {args.manifest}", file=sys.stderr)
        return 1
    if not args.mentions.exists():
        print(f"error: mentions not found: {args.mentions}", file=sys.stderr)
        return 1

    manifest = load_json(args.manifest)
    mentions = load_json(args.mentions)
    drafts = load_json(args.drafts) if args.drafts.exists() else None
    if drafts:
        print(f"including AI drafts from {args.drafts}")

    wb = build_workbook(manifest, mentions, drafts)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    print(f"wrote {args.out} ({len(manifest.get('ias', {}))} IAs, single sheet)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
